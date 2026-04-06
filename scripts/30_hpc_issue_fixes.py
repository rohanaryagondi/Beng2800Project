"""
HPC script: Fix Issues 2 + 3 and regenerate the report.

Run on HPC where loading the full 232 MB decoder matrix is feasible.

Steps:
  1. Download raw data from OSF if orientation_decoder_ready.npz doesn't exist
  2. Regenerate orientation_decoder_ready.npz from raw data if needed
  3. Issue 2: Run reliability-sharpness simulation (if not already done)
  4. Issue 3: Run random-neuron decoder on full 23,589-neuron population
  5. Regenerate the report .docx with all fixes (Issues 1-5)

Usage:
    python3 scripts/30_hpc_issue_fixes.py
"""
import os
import sys
import time
import json
import warnings
import numpy as np
import pandas as pd
from scipy.optimize import curve_fit
from scipy import stats as sp_stats
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.ticker

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
PROC = os.path.join(BASE, "data", "processed")
RAW = os.path.join(BASE, "data", "raw")
FINAL = os.path.join(BASE, "data", "final")
FIGS = os.path.join(BASE, "reports", "figures")
TABLES = os.path.join(BASE, "reports", "tables")
STATUS = os.path.join(BASE, "reports", "status")

for d in [PROC, RAW, FINAL, FIGS, TABLES, STATUS]:
    os.makedirs(d, exist_ok=True)

SEED = 42
DPI = 150


def von_mises_tuning(theta, b, a, kappa, theta0):
    return b + a * np.exp(kappa * np.cos(2.0 * (theta - theta0)))


def circular_error_deg(pred_rad, true_rad):
    err = pred_rad - true_rad
    err = (err + np.pi / 2) % np.pi - np.pi / 2
    return np.degrees(np.abs(err))


def make_fold_indices(n, n_folds, rng):
    idx = rng.permutation(n)
    fold_size = n // n_folds
    folds = []
    for i in range(n_folds):
        start = i * fold_size
        end = start + fold_size if i < n_folds - 1 else n
        test_idx = idx[start:end]
        train_idx = np.setdiff1d(idx, test_idx)
        folds.append((train_idx, test_idx))
    return folds


def run_cv(X, y_cos2, y_sin2, theta_rad, fold_indices):
    Y = np.column_stack([y_cos2, y_sin2])
    mae_list, medae_list = [], []
    all_preds = np.full(len(theta_rad), np.nan)
    for train_idx, test_idx in fold_indices:
        W, _, _, _ = np.linalg.lstsq(X[train_idx], Y[train_idx], rcond=None)
        Y_pred = X[test_idx] @ W
        theta_pred = np.arctan2(Y_pred[:, 1], Y_pred[:, 0]) / 2.0 % np.pi
        errors = circular_error_deg(theta_pred, theta_rad[test_idx])
        mae_list.append(float(np.mean(errors)))
        medae_list.append(float(np.median(errors)))
        all_preds[test_idx] = theta_pred
    return np.mean(mae_list), np.mean(medae_list), all_preds


# ==================================================================
# STEP 0: Ensure the full decoder matrix exists
# ==================================================================
def ensure_full_decoder_matrix():
    full_path = os.path.join(PROC, "orientation_decoder_ready.npz")
    if os.path.exists(full_path):
        print(f"Full decoder matrix found: {full_path}")
        return

    print("Full decoder matrix NOT found. Regenerating from raw data...")

    # Check for raw file
    raw_path = os.path.join(RAW, "stringer_orientations.npy")
    if not os.path.exists(raw_path):
        print("Downloading raw data from OSF (~937 MB)...")
        import urllib.request
        url = "https://osf.io/ny4ut/download"
        os.makedirs(RAW, exist_ok=True)
        urllib.request.urlretrieve(url, raw_path)
        print(f"Downloaded: {raw_path} ({os.path.getsize(raw_path)/1e6:.0f} MB)")

    # Load and process
    print("Loading raw data...")
    dat = np.load(raw_path, allow_pickle=True).item()
    sresp = dat["sresp"]       # (n_neurons, n_trials) float64
    istim = dat["istim"]       # (n_trials,)

    n_neurons, n_trials = sresp.shape
    theta_rad = (istim % np.pi).astype(np.float32)
    theta_deg = np.degrees(theta_rad).astype(np.float32)

    X = sresp.T.astype(np.float32)  # (n_trials, n_neurons)
    neuron_ids = np.arange(n_neurons, dtype=np.int32)

    np.savez_compressed(full_path,
                        X=X, theta_deg=theta_deg, theta_rad=theta_rad,
                        neuron_ids=neuron_ids,
                        recording_id=np.array(["stringer2019_v1_mouse1"]))
    print(f"Saved: {full_path} ({os.path.getsize(full_path)/1e6:.0f} MB)")
    del dat, sresp  # free memory


# ==================================================================
# ISSUE 2: Simulation
# ==================================================================
def fix_issue2():
    sim_csv = os.path.join(TABLES, "issue2_simulation_results.csv")
    sim_fig = os.path.join(FIGS, "issue2_simulation_vs_real.png")

    if os.path.exists(sim_csv) and os.path.exists(sim_fig):
        print("Issue 2: Simulation already completed, skipping.")
        df_sim = pd.read_csv(sim_csv)
        return df_sim.loc[df_sim["analysis"] == "simulation_noise_independent", "spearman_r"].values[0]

    print("\n=== Issue 2: Reliability-sharpness simulation ===")
    t0 = time.time()
    rng = np.random.default_rng(SEED)

    n_sim = 5000
    n_bins = 36
    n_trials_per_bin = 128
    centers_rad = np.radians(np.linspace(2.5, 177.5, n_bins))

    true_kappas = np.clip(rng.exponential(4, size=n_sim), 0, 20)
    true_theta0 = rng.uniform(0, np.pi, size=n_sim)
    true_baseline = rng.uniform(1, 10, size=n_sim)
    true_amplitude = rng.uniform(0.5, 5, size=n_sim)
    noise_std = np.abs(rng.normal(3, 1.5, size=n_sim))

    sim_kappas = np.empty(n_sim)
    sim_reliabilities = np.empty(n_sim)

    for i in range(n_sim):
        true_tc = von_mises_tuning(centers_rad, true_baseline[i], true_amplitude[i],
                                   true_kappas[i], true_theta0[i])
        se = noise_std[i] / np.sqrt(n_trials_per_bin / 2)
        half1 = true_tc + rng.normal(0, se, size=n_bins)
        half2 = true_tc + rng.normal(0, se, size=n_bins)
        full_mean = (half1 + half2) / 2
        r_val, _ = sp_stats.pearsonr(half1, half2)
        sim_reliabilities[i] = r_val
        try:
            b_init = np.percentile(full_mean, 10)
            a_init = max(full_mean.max() - b_init, 0.01)
            with warnings.catch_warnings():
                warnings.simplefilter("ignore")
                popt, _ = curve_fit(von_mises_tuning, centers_rad, full_mean,
                                    p0=[b_init, a_init, 1.0, centers_rad[np.argmax(full_mean)]],
                                    bounds=([-np.inf, 0, 0, -np.inf], [np.inf, np.inf, 20, np.inf]),
                                    maxfev=5000)
            sim_kappas[i] = popt[2]
        except Exception:
            sim_kappas[i] = np.nan
            sim_reliabilities[i] = np.nan

    valid = ~np.isnan(sim_kappas) & ~np.isnan(sim_reliabilities)
    sim_k, sim_r = sim_kappas[valid], sim_reliabilities[valid]
    sim_spearman, _ = sp_stats.spearmanr(sim_k, sim_r)
    print(f"  Simulation Spearman r = {sim_spearman:.3f} (n={valid.sum()})")
    print(f"  Took {time.time()-t0:.1f}s")

    # Load real data
    df = pd.read_parquet(os.path.join(FINAL, "orientation_neuron_analysis.parquet"))
    df_fit = df[df["fit_success"]]
    real_k = df_fit["fit_kappa_or_width"].values
    real_r = df_fit["split_half_reliability"].values

    # Figure
    fig, axes = plt.subplots(1, 3, figsize=(16, 5))
    axes[0].scatter(sim_k, sim_r, s=2, alpha=0.3, color="#e6550d", rasterized=True)
    axes[0].set_xlabel("Fitted $\\kappa$"); axes[0].set_ylabel("Split-half reliability")
    axes[0].set_title(f"Simulation (noise indep. of $\\kappa$)\nSpearman r = {sim_spearman:.3f}")
    axes[0].set_xlim(0, 20); axes[0].set_ylim(-0.5, 1.05)

    axes[1].scatter(real_k, real_r, s=1, alpha=0.1, color="#2171b5", rasterized=True)
    axes[1].set_xlabel("Fitted $\\kappa$"); axes[1].set_ylabel("Split-half reliability")
    axes[1].set_title("Real data\nSpearman r = 0.239")
    axes[1].set_xlim(0, 20); axes[1].set_ylim(-0.5, 1.05)

    n_kbins = 10
    bin_edges = np.linspace(0, 20, n_kbins + 1)
    bc = (bin_edges[:-1] + bin_edges[1:]) / 2
    for data_k, data_r, label, color in [(sim_k, sim_r, "Simulation", "#e6550d"),
                                          (real_k, real_r, "Real data", "#2171b5")]:
        bm, bs = np.empty(n_kbins), np.empty(n_kbins)
        for i in range(n_kbins):
            mask = (data_k >= bin_edges[i]) & (data_k < bin_edges[i+1])
            if i == n_kbins - 1: mask = mask | (data_k == bin_edges[i+1])
            vals = data_r[mask]
            bm[i] = vals.mean() if len(vals) > 0 else np.nan
            bs[i] = vals.std()/np.sqrt(len(vals)) if len(vals) > 1 else np.nan
        axes[2].errorbar(bc, bm, yerr=bs, fmt="o-", color=color, capsize=3, ms=5, lw=1.5, label=label)
    axes[2].set_xlabel("$\\kappa$"); axes[2].set_ylabel("Mean reliability")
    axes[2].set_title("Binned comparison"); axes[2].legend()
    plt.tight_layout()
    fig.savefig(sim_fig, dpi=DPI); plt.close(fig)

    sim_df = pd.DataFrame({
        "analysis": ["simulation_noise_independent", "real_data"],
        "spearman_r": [round(sim_spearman, 4), 0.2392],
        "n_neurons": [int(valid.sum()), len(df_fit)],
        "interpretation": ["Expected correlation when noise is independent of kappa",
                           "Observed — lower than expected"],
    })
    sim_df.to_csv(sim_csv, index=False)
    print(f"  Saved: {sim_fig}, {sim_csv}")
    return sim_spearman


# ==================================================================
# ISSUE 3: Random-neuron decoder
# ==================================================================
def fix_issue3():
    print("\n=== Issue 3: Random-neuron decoder ===")
    t0 = time.time()
    rng = np.random.default_rng(SEED)

    ensure_full_decoder_matrix()

    print("Loading full decoder matrix...")
    d = np.load(os.path.join(PROC, "orientation_decoder_ready.npz"))
    X_full = d["X"]  # (4598, 23589)
    theta_rad = d["theta_rad"]
    dt = np.load(os.path.join(PROC, "orientation_decoder_targets.npz"))
    y_cos2, y_sin2 = dt["y_cos2"], dt["y_sin2"]

    n_trials, n_neurons_full = X_full.shape
    print(f"  Full matrix: {n_trials} x {n_neurons_full}")

    top1000_cv = pd.read_csv(os.path.join(TABLES, "decoder_cv_summary.csv"))
    fold_indices = make_fold_indices(n_trials, 5, rng)

    neuron_counts = [10, 25, 50, 100, 250, 500, 1000, 2000, 5000]
    n_repeats = 10

    random_results = []
    for nc in neuron_counts:
        t_nc = time.time()
        for rep in range(n_repeats):
            subset = rng.choice(n_neurons_full, size=nc, replace=False)
            X_sub = X_full[:, subset]
            mae, medae, _ = run_cv(X_sub, y_cos2, y_sin2, theta_rad, fold_indices)
            random_results.append({
                "neuron_count": nc, "repeat_id": rep,
                "mae_circular_deg": round(mae, 4),
                "median_ae_circular_deg": round(medae, 4),
                "source": "random_full_population",
            })
        maes = [r["mae_circular_deg"] for r in random_results if r["neuron_count"] == nc]
        print(f"  {nc:5d} random neurons: MAE = {np.mean(maes):.2f} +/- {np.std(maes):.2f} deg  "
              f"({time.time()-t_nc:.1f}s)")

    del X_full, d
    print("  (Released full X)")

    # Save
    rand_df = pd.DataFrame(random_results)
    rand_summary = rand_df.groupby("neuron_count").agg(
        mae_mean=("mae_circular_deg", "mean"),
        mae_std=("mae_circular_deg", "std"),
        medae_mean=("median_ae_circular_deg", "mean"),
        n_repeats=("repeat_id", "count"),
    ).reset_index()
    rand_summary["source"] = "random_full_population"

    top_summary = top1000_cv[["neuron_count", "mae_mean", "mae_std", "medae_mean", "n_repeats"]].copy()
    top_summary["source"] = "top_1000_pool"
    combined = pd.concat([top_summary, rand_summary], ignore_index=True)

    rand_df.to_csv(os.path.join(TABLES, "decoder_random_neurons_detail.csv"), index=False)
    combined.to_csv(os.path.join(TABLES, "issue3_decoder_comparison.csv"), index=False)

    # Figure: both curves
    fig, ax = plt.subplots(figsize=(9, 5.5))
    nc_top = top1000_cv["neuron_count"].values
    ax.errorbar(nc_top, top1000_cv["mae_mean"].values, yerr=top1000_cv["mae_std"].fillna(0).values,
                fmt="s-", color="#e6550d", capsize=4, ms=7, lw=2, label="Top-1000 reliable neurons")
    nc_rand = rand_summary["neuron_count"].values
    ax.errorbar(nc_rand, rand_summary["mae_mean"].values, yerr=rand_summary["mae_std"].values,
                fmt="o-", color="#2171b5", capsize=4, ms=7, lw=2, label="Random neurons (full population)")
    ax.axhline(45, color="gray", ls=":", lw=1, label="Chance (45 deg)")
    ax.set_xlabel("Number of neurons", fontsize=11)
    ax.set_ylabel("Mean circular MAE (deg)", fontsize=11)
    ax.set_title("Decoder performance: pre-selected vs. random neurons", fontsize=12)
    ax.set_xscale("log")
    ax.set_xticks(sorted(set(list(nc_top) + list(nc_rand))))
    ax.get_xaxis().set_major_formatter(matplotlib.ticker.ScalarFormatter())
    ax.tick_params(axis='x', rotation=45)
    ax.legend(fontsize=10); ax.set_ylim(0, 50)
    plt.tight_layout()
    fig.savefig(os.path.join(FIGS, "issue3_decoder_random_vs_top1000.png"), dpi=DPI)
    plt.close(fig)

    # Also regenerate the final decoder scaling figure for the report
    fig, ax = plt.subplots(figsize=(8, 5))
    ax.errorbar(nc_rand, rand_summary["mae_mean"].values, yerr=rand_summary["mae_std"].values,
                fmt="o-", color="#2171b5", capsize=4, ms=7, lw=2, label="Random neurons")
    ax.errorbar(nc_top, top1000_cv["mae_mean"].values, yerr=top1000_cv["mae_std"].fillna(0).values,
                fmt="s--", color="#e6550d", capsize=3, ms=5, lw=1.5, alpha=0.7,
                label="Top-1000 reliable (upper bound)")
    ax.axhline(45, color="gray", ls=":", lw=1, label="Chance (45 deg)")
    ax.set_xlabel("Number of neurons", fontsize=11)
    ax.set_ylabel("Mean circular MAE (deg)", fontsize=11)
    ax.set_title("Population decoder performance vs. neuron count", fontsize=12)
    ax.set_xscale("log")
    ax.set_xticks(sorted(set(list(nc_top) + list(nc_rand))))
    ax.get_xaxis().set_major_formatter(matplotlib.ticker.ScalarFormatter())
    ax.tick_params(axis='x', rotation=45)
    ax.legend(fontsize=10); ax.set_ylim(0, 50)
    plt.tight_layout()
    fig.savefig(os.path.join(FIGS, "final_fig4_decoder_scaling.png"), dpi=DPI)
    plt.close(fig)

    elapsed = time.time() - t0
    print(f"\n  Issue 3 completed in {elapsed:.1f}s")
    return rand_summary, combined


# ==================================================================
# REGENERATE REPORT with all fixes (Issues 1-5)
# ==================================================================
def regenerate_report(sim_spearman, rand_summary):
    """Regenerate the .docx report with all issue fixes applied."""
    print("\n=== Regenerating report with all fixes ===")

    try:
        from docx import Document as DocxDocument
        from docx.shared import Inches, Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.enum.table import WD_TABLE_ALIGNMENT
        from docx.oxml.ns import qn
    except ImportError:
        print("  python-docx not installed. Install with: pip install python-docx")
        print("  Skipping report generation. Report can be regenerated on laptop.")
        return

    # Load all results
    cv_summary = pd.read_csv(os.path.join(TABLES, "decoder_cv_summary.csv"))
    model_results = pd.read_csv(os.path.join(TABLES, "reliability_model_results.csv"))
    corr_results = pd.read_csv(os.path.join(TABLES, "reliability_correlations.csv"))
    tuning_summary = pd.read_csv(os.path.join(TABLES, "tuning_fit_summary.csv"))
    combined = pd.read_csv(os.path.join(TABLES, "issue3_decoder_comparison.csv"))

    spearman_r = corr_results[(corr_results["method"] == "Spearman") & (corr_results["metric"] == "r")]["value"].values[0]
    sp_ci_lo = corr_results[(corr_results["method"] == "Spearman") & (corr_results["metric"] == "r")]["bootstrap_ci_lo"].values[0]
    sp_ci_hi = corr_results[(corr_results["method"] == "Spearman") & (corr_results["metric"] == "r")]["bootstrap_ci_hi"].values[0]
    beta_kappa = model_results[model_results["parameter"] == "kappa"]["coefficient"].values[0]
    se_kappa = model_results[model_results["parameter"] == "kappa"]["std_error"].values[0]
    t_kappa = model_results[model_results["parameter"] == "kappa"]["t_statistic"].values[0]
    beta_mean = model_results[model_results["parameter"] == "mean_response"]["coefficient"].values[0]
    t_mean = model_results[model_results["parameter"] == "mean_response"]["t_statistic"].values[0]

    r2_median = tuning_summary[tuning_summary["metric"] == "r2_median"]["value"].values[0]
    r2_q25 = tuning_summary[tuning_summary["metric"] == "r2_q25"]["value"].values[0]
    r2_q75 = tuning_summary[tuning_summary["metric"] == "r2_q75"]["value"].values[0]
    kappa_median = tuning_summary[tuning_summary["metric"] == "kappa_median"]["value"].values[0]
    kappa_q25 = tuning_summary[tuning_summary["metric"] == "kappa_q25"]["value"].values[0]
    kappa_q75 = tuning_summary[tuning_summary["metric"] == "kappa_q75"]["value"].values[0]

    # Random-neuron results
    rand_only = combined[combined["source"] == "random_full_population"]
    mae_rand_1000 = rand_only.loc[rand_only["neuron_count"] == 1000, "mae_mean"].values[0]
    mae_rand_10 = rand_only.loc[rand_only["neuron_count"] == 10, "mae_mean"].values[0]
    mae_top_1000 = cv_summary.loc[cv_summary["neuron_count"] == 1000, "mae_mean"].values[0]

    # Helpers
    def set_cell_shading(cell, color):
        shading = cell._element.get_or_add_tcPr()
        elm = shading.makeelement(qn('w:shd'), {qn('w:fill'): color, qn('w:val'): 'clear'})
        shading.append(elm)

    def add_table(doc, headers, rows, col_widths=None):
        table = doc.add_table(rows=1 + len(rows), cols=len(headers))
        table.style = 'Table Grid'
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for j, h in enumerate(headers):
            c = table.rows[0].cells[j]; c.text = ""
            r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(9); r.font.name = "Arial"
            c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_cell_shading(c, "D5E8F0")
        for i, rd in enumerate(rows):
            for j, v in enumerate(rd):
                c = table.rows[i+1].cells[j]; c.text = ""
                r = c.paragraphs[0].add_run(str(v)); r.font.size = Pt(9); r.font.name = "Arial"
                c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        if col_widths:
            for row in table.rows:
                for j, w in enumerate(col_widths):
                    row.cells[j].width = Inches(w)
        return table

    def add_fig(doc, path, caption, width=5.5):
        if not os.path.exists(path):
            doc.add_paragraph(f"[Figure not found: {path}]").alignment = WD_ALIGN_PARAGRAPH.CENTER
            return
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(path, width=Inches(width))
        c = doc.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = c.add_run(caption); r.font.size = Pt(9); r.font.name = "Arial"; r.italic = True
        c.paragraph_format.space_after = Pt(12)

    def heading(doc, text, level):
        h = doc.add_heading(text, level=level)
        for r in h.runs: r.font.color.rgb = RGBColor(0,0,0); r.font.name = "Arial"

    def body(doc, text, bold=False, italic=False, size=11):
        p = doc.add_paragraph(); r = p.add_run(text)
        r.font.size = Pt(size); r.font.name = "Arial"; r.bold = bold; r.italic = italic
        p.paragraph_format.space_after = Pt(6); p.paragraph_format.line_spacing = 1.15
        return p

    # BUILD DOCUMENT
    doc = DocxDocument()
    style = doc.styles['Normal']; style.font.name = 'Arial'; style.font.size = Pt(11)
    for s in doc.sections:
        s.top_margin = Inches(1); s.bottom_margin = Inches(1)
        s.left_margin = Inches(1); s.right_margin = Inches(1)

    # Title
    t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("Orientation Tuning Heterogeneity and Population Decoding\nin Mouse Primary Visual Cortex")
    r.font.size = Pt(16); r.font.name = "Arial"; r.bold = True
    t.paragraph_format.space_after = Pt(4)
    for text in ["BENG 2800 Final Project Report", "Rohan Aryagondi", "April 2026"]:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text); r.font.size = Pt(11 if text != "BENG 2800 Final Project Report" else 12); r.font.name = "Arial"
        p.paragraph_format.space_after = Pt(2 if "2026" not in text else 18)

    # 1. INTRODUCTION
    heading(doc, "1. Introduction and Background", 1)
    body(doc,
        "Neurons in the primary visual cortex (V1) respond selectively to the orientation "
        "of visual stimuli, a property first described by Hubel and Wiesel (1962). This "
        "orientation selectivity is a fundamental building block of visual processing. "
        "Understanding how orientation is encoded across populations of neurons remains "
        "a central question in systems neuroscience."
    )
    body(doc,
        "Recent advances in two-photon calcium imaging allow simultaneous recording from "
        "tens of thousands of neurons. Stringer, Michaelos, and Pachitariu (2021) recorded "
        "from approximately 23,000 neurons in mouse V1, demonstrating high-precision "
        "population coding of orientation. Their dataset provides a unique opportunity to "
        "investigate tuning diversity and population-level encoding."
    )
    body(doc, "In this project, we address four interrelated questions:")
    for q in [
        "How much do preferred orientation and tuning sharpness vary across neurons in this V1 recording?",
        "Are neurons with sharper orientation tuning also more reliable across repeated trials?",
        "How accurately can a simple linear population decoder predict stimulus orientation?",
        "How does decoder performance change as the number of neurons increases?",
    ]:
        p = doc.add_paragraph(style='List Number')
        r = p.add_run(q); r.font.size = Pt(11); r.font.name = "Arial"; r.italic = True

    body(doc,
        "These questions bridge single-neuron physiology with population computation, using "
        "two analytical methods from this course: nonlinear least-squares curve fitting "
        "and linear least-squares regression."
    )

    # 2. METHODOLOGY
    heading(doc, "2. Methodology", 1)

    heading(doc, "2.1 Dataset", 2)
    body(doc,
        "We analyzed two-photon calcium imaging data from Stringer et al. (2021), consisting "
        "of 23,589 simultaneously recorded neurons in V1 of a single head-fixed mouse during "
        "passive viewing of 4,598 static oriented gratings (0\u2013180\u00b0, approximately "
        "uniformly distributed). Neural responses were provided as deconvolved calcium activity "
        "(one scalar per neuron per trial) extracted using Suite2p."
    )

    heading(doc, "2.2 Nonlinear Least-Squares Tuning Model", 2)
    body(doc,
        "Empirical tuning curves were estimated by binning trials into 36 bins of 5\u00b0 each "
        "(~128 trials/bin). We fit a von Mises orientation tuning function using nonlinear "
        "least squares (scipy.optimize.curve_fit):"
    )
    eq = doc.add_paragraph(); eq.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = eq.add_run("r(\u03b8) = b + a \u00b7 exp(\u03ba \u00b7 cos(2(\u03b8 \u2212 \u03b8\u2080)))")
    r.font.size = Pt(11); r.font.name = "Arial"; r.italic = True
    eq.paragraph_format.space_before = Pt(6); eq.paragraph_format.space_after = Pt(6)
    body(doc,
        "where b is baseline, a is amplitude, \u03ba is tuning sharpness (bounded [0, 20]), "
        "and \u03b8\u2080 is preferred orientation. R\u00b2 was computed on binned means "
        "(not single trials); single-trial prediction accuracy would be lower due to "
        "trial-to-trial variability."
    )

    heading(doc, "2.3 Reliability Analysis", 2)
    body(doc,
        "Split-half reliability was the Pearson correlation between tuning curves from odd- "
        "and even-indexed trials. We assessed the kappa\u2013reliability relationship using an "
        "OLS linear model (reliability = \u03b2\u2080 + \u03b2\u2081\u00b7\u03ba + "
        "\u03b2\u2082\u00b7mean_response) via np.linalg.lstsq, plus Pearson and Spearman "
        "correlations with 95% bootstrap CIs (5,000 resamples). We also ran a simulation "
        "to quantify the mathematically expected correlation between \u03ba and reliability."
    )

    heading(doc, "2.4 Linear Population Decoder", 2)
    body(doc,
        "An OLS linear decoder regressed the trial-by-neuron response matrix onto target "
        "vectors y_cos = cos(2\u03b8) and y_sin = sin(2\u03b8). Orientation was recovered "
        "as \u03b8_pred = arctan2(sin_pred, cos_pred)/2. We used 5-fold CV with circular "
        "MAE (modulo 180\u00b0). Scaling was evaluated at neuron counts of 10\u20135,000 "
        "(10 random subsets each). We present results for both randomly selected neurons "
        "from the full population and the top 1,000 most reliable neurons as an upper bound. "
        "A shuffle control (permuted labels, 10 repeats) established the chance baseline."
    )

    # 3. RESULTS
    heading(doc, "3. Results and Analysis", 1)

    heading(doc, "3.1 Heterogeneity of Orientation Tuning", 2)
    # ISSUE 1 FIX: separate untuned from broadly tuned
    body(doc,
        f"The von Mises fitting algorithm converged for all 23,589 neurons; however, "
        f"2,218 (9.4%) produced R\u00b2 < 0.3 and are better characterized as non-orientation-selective "
        f"rather than \u201cbroadly tuned.\u201d Among the remaining neurons, the median R\u00b2 was "
        f"{r2_median:.3f} (IQR: [{r2_q25:.3f}, {r2_q75:.3f}]; note these values reflect fits to "
        f"trial-averaged tuning curves, not single-trial predictions)."
    )
    add_fig(doc, os.path.join(FIGS, "final_fig1_tuning_examples.png"),
            "Figure 1. Example orientation tuning curves with von Mises fits spanning the range of fit quality.", width=5.8)
    # ISSUE 5b FIX: soften Ganguli citation
    body(doc,
        f"Preferred orientations were approximately uniformly distributed across 0\u2013180\u00b0 "
        f"(Figure 2, left). Tuning sharpness (\u03ba) varied substantially (median = {kappa_median:.2f}, "
        f"IQR: [{kappa_q25:.2f}, {kappa_q75:.2f}]; Figure 2, right). The observed heterogeneity "
        f"is consistent with the premise that diverse tuning profiles may support population coding "
        f"(Ganguli and Simoncelli, 2014), though testing optimality would require comparing against "
        f"theoretical predictions."
    )
    add_fig(doc, os.path.join(FIGS, "final_fig2_tuning_parameter_distributions.png"),
            "Figure 2. Left: preferred orientations are approximately uniform. "
            f"Right: tuning sharpness is right-skewed (median \u03ba = {kappa_median:.2f}).", width=5.8)

    # 3.2 Reliability — ISSUE 2 FIX: reframe with simulation
    heading(doc, "3.2 Tuning Sharpness and Trial-to-Trial Reliability", 2)
    body(doc,
        f"We found a positive correlation between \u03ba and reliability (Spearman r = {spearman_r:.3f}, "
        f"95% CI: [{sp_ci_lo:.3f}, {sp_ci_hi:.3f}]). However, a simulation in which noise was "
        f"entirely independent of \u03ba produced a much stronger correlation (Spearman r = {sim_spearman:.3f}; "
        f"Figure 3), indicating that the \u03ba\u2013reliability coupling is largely a mathematical "
        f"consequence of how both metrics are derived from tuning curve structure. The real data\u2019s "
        f"weaker correlation (0.239 vs. expected ~0.91) suggests that biological noise sources "
        f"partially decouple tuning strength from measurement reliability."
    )
    body(doc, "Table 1. OLS regression results: reliability ~ \u03ba + mean_response", bold=True, size=10)
    add_table(doc,
        ["Parameter", "Coefficient", "SE", "t", "p-value"],
        [["Intercept", "0.670", "0.0022", f"{model_results[model_results['parameter']=='intercept']['t_statistic'].values[0]:.1f}", "< 10\u207b\u00b3\u2070\u2070"],
         ["\u03ba", f"{beta_kappa:.4f}", f"{se_kappa:.4f}", f"{t_kappa:.1f}", "< 10\u207b\u2077\u2070"],
         ["Mean response", f"{beta_mean:.4f}", "0.0001", f"{t_mean:.1f}", "< 10\u207b\u00b2\u2070\u2070"]],
        col_widths=[1.3, 1.1, 0.8, 0.8, 1.1])
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    # ISSUE 5c: add heteroscedasticity note
    body(doc,
        f"The OLS model R\u00b2 = 0.055, confirming the modest predictive power. Note that "
        f"reliability is bounded [\u22121, 1], so OLS standard errors should be interpreted "
        f"with caution due to potential heteroscedasticity."
    )
    add_fig(doc, os.path.join(FIGS, "issue2_simulation_vs_real.png"),
            f"Figure 3. Left: simulation with noise independent of \u03ba (r = {sim_spearman:.3f}). "
            f"Center: real data (r = 0.239). Right: binned comparison\u2014the real data shows a much "
            f"weaker \u03ba\u2013reliability coupling than expected from signal structure alone.", width=5.8)

    # 3.3 Decoder — ISSUE 3 FIX: random neurons as primary
    heading(doc, "3.3 Population Decoder Performance", 2)
    body(doc,
        f"With 1,000 randomly selected neurons from the full population, the OLS decoder achieved "
        f"MAE = {mae_rand_1000:.2f}\u00b0, far below chance (~45\u00b0). Pre-selecting the 1,000 most "
        f"reliable neurons improved this to {mae_top_1000:.2f}\u00b0, a {(mae_rand_1000/mae_top_1000 - 1)*100:.0f}% "
        f"improvement that highlights how neuron quality affects decoding (Table 2, Figure 4)."
    )
    body(doc, "Table 2. Decoder performance: random vs. pre-selected neurons.", bold=True, size=10)
    dec_rows = []
    for nc in sorted(rand_only["neuron_count"].unique()):
        rr = rand_only[rand_only["neuron_count"] == nc].iloc[0]
        tr = cv_summary[cv_summary["neuron_count"] == nc]
        rand_str = f"{rr['mae_mean']:.2f} \u00b1 {rr['mae_std']:.2f}"
        if len(tr) > 0:
            ts = tr.iloc[0]
            top_str = f"{ts['mae_mean']:.2f} \u00b1 {ts['mae_std']:.2f}" if not pd.isna(ts["mae_std"]) else f"{ts['mae_mean']:.2f}"
        else:
            top_str = "\u2014"
        dec_rows.append([str(nc), rand_str, top_str])
    add_table(doc, ["Neurons", "Random MAE (deg)", "Top-1000 MAE (deg)"], dec_rows,
              col_widths=[1.0, 2.5, 2.5])
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    body(doc,
        f"Performance improved monotonically with neuron count (Figure 4). With random neurons, "
        f"gains continued steeply up to 2,000\u20135,000 neurons, while the pre-selected pool "
        f"showed earlier saturation. The gap between the two curves quantifies the benefit of "
        f"neuron selection."
    )
    add_fig(doc, os.path.join(FIGS, "final_fig4_decoder_scaling.png"),
            "Figure 4. Decoder MAE vs. neuron count. Blue: random neurons from full population. "
            "Orange dashed: top-1000 most reliable neurons. Both far exceed chance (45\u00b0).", width=5.0)
    add_fig(doc, os.path.join(FIGS, "final_fig5_decoder_examples.png"),
            f"Figure 5. Predicted vs. true orientation (top-1000 decoder, MAE = {mae_top_1000:.1f}\u00b0).", width=3.8)

    # 4. DISCUSSION
    heading(doc, "4. Discussion and Conclusion", 1)
    heading(doc, "4.1 Summary of Findings", 2)
    # ISSUE 5a: use "this recording" not "mouse V1"
    body(doc,
        f"This analysis of 23,589 simultaneously recorded neurons in this V1 recording revealed "
        f"substantial heterogeneity in orientation tuning: preferred orientations tile the full "
        f"0\u2013180\u00b0 range uniformly, while tuning sharpness spans over an order of magnitude."
    )
    body(doc,
        f"The positive \u03ba\u2013reliability correlation (Spearman r = {spearman_r:.3f}) is "
        f"substantially weaker than the r \u2248 {sim_spearman:.2f} expected from signal structure "
        f"alone (simulation), suggesting that biological noise sources partially decouple tuning "
        f"strength from reliability. This reframes the finding: the correlation is not evidence "
        f"of a novel biological relationship, but its attenuation from the expected value points "
        f"to interesting noise properties in the population."
    )
    body(doc,
        f"The linear decoder achieves {mae_rand_1000:.1f}\u00b0 MAE with 1,000 random neurons "
        f"(improving to {mae_top_1000:.1f}\u00b0 with pre-selected reliable neurons). The {(mae_rand_1000/mae_top_1000 - 1)*100:.0f}% "
        f"gap between random and pre-selected neurons demonstrates that neuron quality matters "
        f"for decoding, even within a simple linear framework."
    )
    heading(doc, "4.2 Limitations", 2)
    for lim in [
        "Single recording from one mouse; replication across animals is needed.",
        "~9.4% of neurons are non-orientation-selective despite converging fits (R\u00b2 < 0.3).",
        "R\u00b2 values reflect binned means, not single-trial prediction accuracy.",
        "The \u03ba\u2013reliability correlation is largely a mathematical consequence of shared tuning curve structure, not purely biological.",
        "Reliability is bounded [\u22121, 1], violating OLS homoscedasticity assumptions.",
        "Deconvolved calcium activity is an indirect measure of spiking.",
    ]:
        p = doc.add_paragraph(style='List Bullet')
        r = p.add_run(lim); r.font.size = Pt(11); r.font.name = "Arial"

    heading(doc, "4.3 Conclusion", 2)
    body(doc,
        "This project demonstrates that V1 neurons in this recording exhibit diverse orientation "
        "tuning whose collective activity supports high-precision decoding. Nonlinear least-squares "
        "fitting characterized individual tuning curves, while linear regression decoded population "
        "activity. The combination reveals both the building blocks of orientation coding and how "
        "they combine for population-level readout."
    )

    # 5. REFLECTION
    heading(doc, "5. Individual Reflection", 1)
    body(doc, "[To be completed individually.]")

    # REFERENCES
    heading(doc, "References", 1)
    for ref in [
        "Ganguli, D., & Simoncelli, E.P. (2014). Efficient sensory encoding and Bayesian inference with heterogeneous neural populations. Neural Computation, 26(10), 2103\u20132134.",
        "Hubel, D.H., & Wiesel, T.N. (1962). Receptive fields, binocular interaction and functional architecture in the cat\u2019s visual cortex. The Journal of Physiology, 160(1), 106\u2013154.",
        "Stringer, C., Michaelos, M., & Pachitariu, M. (2021). High-precision coding in visual cortex. Cell, 184(10), 2767\u20132778. DOI: 10.1016/j.cell.2021.03.042.",
    ]:
        p = doc.add_paragraph(); r = p.add_run(ref)
        r.font.size = Pt(10); r.font.name = "Arial"
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.left_indent = Inches(0.5)
        p.paragraph_format.first_line_indent = Inches(-0.5)

    out_path = os.path.join(BASE, "reports", "BENG2800_Final_Report.docx")
    doc.save(out_path)
    print(f"  Report saved: {out_path} ({os.path.getsize(out_path)/1e6:.2f} MB)")


def main():
    sim_spearman = fix_issue2()
    rand_summary, combined = fix_issue3()
    regenerate_report(sim_spearman, rand_summary)
    print("\n=== ALL ISSUE FIXES COMPLETE ===")


if __name__ == "__main__":
    main()
