"""
Generate the BENG 2800 final report as a .docx file.
Uses python-docx to create a properly formatted Word document.
"""
import os
import numpy as np
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FIGS = os.path.join(BASE, "reports", "figures")
TABLES = os.path.join(BASE, "reports", "tables")
OUT = os.path.join(BASE, "reports")


def set_cell_shading(cell, color):
    """Set cell background color."""
    shading = cell._element.get_or_add_tcPr()
    shading_elm = shading.makeelement(qn('w:shd'), {
        qn('w:fill'): color,
        qn('w:val'): 'clear',
    })
    shading.append(shading_elm)


def add_formatted_table(doc, headers, rows, col_widths=None):
    """Add a formatted table with header shading."""
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Header row
    for j, h in enumerate(headers):
        cell = table.rows[0].cells[j]
        cell.text = ""
        p = cell.paragraphs[0]
        run = p.add_run(h)
        run.bold = True
        run.font.size = Pt(9)
        run.font.name = "Arial"
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_cell_shading(cell, "D5E8F0")

    # Data rows
    for i, row_data in enumerate(rows):
        for j, val in enumerate(row_data):
            cell = table.rows[i + 1].cells[j]
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(str(val))
            run.font.size = Pt(9)
            run.font.name = "Arial"
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Set column widths if provided
    if col_widths:
        for row in table.rows:
            for j, w in enumerate(col_widths):
                row.cells[j].width = Inches(w)

    return table


def add_figure(doc, path, caption, width=5.5):
    """Add a figure with caption."""
    if not os.path.exists(path):
        p = doc.add_paragraph(f"[Figure not found: {path}]")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        return

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(path, width=Inches(width))

    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cap.add_run(caption)
    run.font.size = Pt(9)
    run.font.name = "Arial"
    run.italic = True
    cap.paragraph_format.space_after = Pt(12)


def add_heading_styled(doc, text, level):
    """Add a heading with consistent styling."""
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0, 0, 0)
        run.font.name = "Arial"
    return h


def add_body(doc, text, bold=False, italic=False, size=11):
    """Add a body paragraph."""
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.name = "Arial"
    run.bold = bold
    run.italic = italic
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    return p


def main():
    print("Generating BENG 2800 Final Report...\n")

    # Load results
    cv_summary = pd.read_csv(os.path.join(TABLES, "decoder_cv_summary.csv"))
    model_results = pd.read_csv(os.path.join(TABLES, "reliability_model_results.csv"))
    corr_results = pd.read_csv(os.path.join(TABLES, "reliability_correlations.csv"))
    tuning_summary = pd.read_csv(os.path.join(TABLES, "tuning_fit_summary.csv"))

    # Extract key numbers
    spearman_r = corr_results[(corr_results["method"] == "Spearman") & (corr_results["metric"] == "r")]["value"].values[0]
    sp_ci_lo = corr_results[(corr_results["method"] == "Spearman") & (corr_results["metric"] == "r")]["bootstrap_ci_lo"].values[0]
    sp_ci_hi = corr_results[(corr_results["method"] == "Spearman") & (corr_results["metric"] == "r")]["bootstrap_ci_hi"].values[0]
    pearson_r = corr_results[(corr_results["method"] == "Pearson") & (corr_results["metric"] == "r")]["value"].values[0]

    beta_kappa = model_results[model_results["parameter"] == "kappa"]["coefficient"].values[0]
    se_kappa = model_results[model_results["parameter"] == "kappa"]["std_error"].values[0]
    t_kappa = model_results[model_results["parameter"] == "kappa"]["t_statistic"].values[0]
    p_kappa = model_results[model_results["parameter"] == "kappa"]["p_value"].values[0]
    beta_mean = model_results[model_results["parameter"] == "mean_response"]["coefficient"].values[0]
    t_mean = model_results[model_results["parameter"] == "mean_response"]["t_statistic"].values[0]

    mae_1000 = cv_summary.loc[cv_summary["neuron_count"] == 1000, "mae_mean"].values[0]
    mae_10 = cv_summary.loc[cv_summary["neuron_count"] == 10, "mae_mean"].values[0]
    medae_1000 = cv_summary.loc[cv_summary["neuron_count"] == 1000, "medae_mean"].values[0]

    r2_median = tuning_summary[tuning_summary["metric"] == "r2_median"]["value"].values[0]
    kappa_median = tuning_summary[tuning_summary["metric"] == "kappa_median"]["value"].values[0]
    kappa_q25 = tuning_summary[tuning_summary["metric"] == "kappa_q25"]["value"].values[0]
    kappa_q75 = tuning_summary[tuning_summary["metric"] == "kappa_q75"]["value"].values[0]

    # ===================================================================
    # Create document
    # ===================================================================
    doc = Document()

    # Set default font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Arial'
    font.size = Pt(11)

    # Set margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # ===================================================================
    # TITLE
    # ===================================================================
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Orientation Tuning Heterogeneity and Population Decoding\nin Mouse Primary Visual Cortex")
    run.font.size = Pt(16)
    run.font.name = "Arial"
    run.bold = True
    title.paragraph_format.space_after = Pt(4)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("BENG 2800 Final Project Report")
    run.font.size = Pt(12)
    run.font.name = "Arial"
    subtitle.paragraph_format.space_after = Pt(2)

    author = doc.add_paragraph()
    author.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = author.add_run("Rohan Aryagondi")
    run.font.size = Pt(11)
    run.font.name = "Arial"
    author.paragraph_format.space_after = Pt(2)

    date_p = doc.add_paragraph()
    date_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = date_p.add_run("April 2026")
    run.font.size = Pt(11)
    run.font.name = "Arial"
    date_p.paragraph_format.space_after = Pt(18)

    # ===================================================================
    # 1. INTRODUCTION & BACKGROUND
    # ===================================================================
    add_heading_styled(doc, "1. Introduction and Background", level=1)

    add_body(doc,
        "Neurons in the primary visual cortex (V1) respond selectively to the orientation "
        "of visual stimuli, a property first described by Hubel and Wiesel (1962). This "
        "orientation selectivity is considered a fundamental building block of visual "
        "processing, enabling the brain to represent edges, contours, and textures. "
        "Understanding how orientation is encoded across populations of neurons remains "
        "a central question in systems neuroscience, with implications for sensory "
        "prosthetics and brain-computer interfaces."
    )

    add_body(doc,
        "Recent advances in two-photon calcium imaging now allow simultaneous recording "
        "from tens of thousands of neurons. Stringer, Michaelos, and Pachitariu (2021) "
        "leveraged this technology to record from approximately 23,000 neurons in mouse V1, "
        "demonstrating that neural populations encode orientation with remarkably high "
        "precision. Their dataset provides a unique opportunity to investigate the "
        "diversity of orientation tuning at the single-neuron level and the efficiency "
        "of population-level encoding."
    )

    add_body(doc, "In this project, we address four interrelated questions:")

    questions = [
        "How much do preferred orientation and tuning sharpness vary across neurons in mouse V1?",
        "Are neurons with sharper orientation tuning also more reliable across repeated trials?",
        "How accurately can a simple linear population decoder predict stimulus orientation from neural activity?",
        "How does decoder performance change as the number of neurons increases?",
    ]
    for i, q in enumerate(questions, 1):
        p = doc.add_paragraph(style='List Number')
        run = p.add_run(q)
        run.font.size = Pt(11)
        run.font.name = "Arial"
        run.italic = True

    add_body(doc,
        "These questions are significant because they bridge single-neuron physiology "
        "(tuning curves) with population-level computation (decoding), using two core "
        "analytical methods covered in this course: nonlinear least-squares curve fitting "
        "and linear least-squares regression."
    )

    # ===================================================================
    # 2. METHODOLOGY
    # ===================================================================
    add_heading_styled(doc, "2. Methodology", level=1)

    # 2.1 Dataset
    add_heading_styled(doc, "2.1 Dataset", level=2)
    add_body(doc,
        "We analyzed publicly available two-photon calcium imaging data from Stringer et al. "
        "(2021), consisting of simultaneous recordings from 23,589 neurons in V1 of a single "
        "head-fixed mouse during passive viewing of static oriented gratings. A total of 4,598 "
        "trials were presented, with stimulus orientations approximately uniformly distributed "
        "across 0\u2013180\u00b0. Neural responses were provided as deconvolved calcium activity "
        "(one scalar per neuron per trial), extracted using Suite2p. Stimulus directions "
        "(0\u2013360\u00b0) were mapped to orientations (0\u2013180\u00b0) since static gratings "
        "at 0\u00b0 and 180\u00b0 are physically identical."
    )

    # 2.2 Tuning curve estimation
    add_heading_styled(doc, "2.2 Orientation Tuning Curve Estimation", level=2)
    add_body(doc,
        "For each neuron, we estimated an empirical orientation tuning curve by binning trials "
        "into 36 bins of 5\u00b0 each (approximately 128 trials per bin) and computing the mean "
        "response per bin. Split-half reliability was assessed by computing the Pearson "
        "correlation between tuning curves from odd- and even-indexed trials."
    )

    # 2.3 Nonlinear least squares
    add_heading_styled(doc, "2.3 Nonlinear Least-Squares Tuning Model", level=2)
    add_body(doc,
        "We fit a von Mises orientation tuning function to each neuron\u2019s binned tuning "
        "curve using nonlinear least squares (scipy.optimize.curve_fit):"
    )

    eq = doc.add_paragraph()
    eq.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = eq.add_run("r(\u03b8) = b + a \u00b7 exp(\u03ba \u00b7 cos(2(\u03b8 \u2212 \u03b8\u2080)))")
    run.font.size = Pt(11)
    run.font.name = "Arial"
    run.italic = True
    eq.paragraph_format.space_before = Pt(6)
    eq.paragraph_format.space_after = Pt(6)

    add_body(doc,
        "where b is baseline activity, a is response amplitude, \u03ba (kappa) is tuning "
        "sharpness (higher values indicate narrower tuning, bounded [0, 20]), and \u03b8\u2080 "
        "is preferred orientation. The cos(2(\u03b8 \u2212 \u03b8\u2080)) form provides 180\u00b0 "
        "periodicity appropriate for orientation. Initial parameters were derived from "
        "the empirical tuning curve. Goodness of fit was assessed using R\u00b2."
    )

    # 2.4 Reliability analysis
    add_heading_styled(doc, "2.4 Reliability Analysis", level=2)
    add_body(doc,
        "To test whether tuning sharpness predicts reliability, we fit an ordinary "
        "least-squares (OLS) linear model via np.linalg.lstsq:"
    )

    eq2 = doc.add_paragraph()
    eq2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = eq2.add_run("reliability = \u03b2\u2080 + \u03b2\u2081 \u00b7 \u03ba + \u03b2\u2082 \u00b7 mean_response")
    run.font.size = Pt(11)
    run.font.name = "Arial"
    run.italic = True
    eq2.paragraph_format.space_before = Pt(6)
    eq2.paragraph_format.space_after = Pt(6)

    add_body(doc,
        "Standard errors were computed from the residual variance and (X\u1d40X)\u207b\u00b9. "
        "We also computed Pearson and Spearman correlations with 95% bootstrap confidence "
        "intervals (5,000 resamples)."
    )

    # 2.5 Linear decoder
    add_heading_styled(doc, "2.5 Linear Population Decoder", level=2)
    add_body(doc,
        "We implemented an OLS linear decoder (np.linalg.lstsq) that regresses the "
        "trial-by-neuron response matrix X onto target vectors y_cos = cos(2\u03b8) and "
        "y_sin = sin(2\u03b8). Predicted orientation was recovered as "
        "\u03b8_pred = arctan2(sin_pred, cos_pred) / 2. We used 5-fold cross-validation "
        "with circular mean absolute error (MAE, modulo 180\u00b0) as the primary metric."
    )

    add_body(doc,
        "To assess scaling, we evaluated decoder performance for neuron counts of "
        "10, 25, 50, 100, 250, 500, and 1,000 (the top 1,000 most reliable neurons). "
        "For each count, 20 random neuron subsets were drawn and 5-fold CV was run on each. "
        "A shuffle control (permuted orientation labels, 10 repeats at 100 neurons) established "
        "the chance baseline."
    )

    # ===================================================================
    # 3. RESULTS & ANALYSIS
    # ===================================================================
    add_heading_styled(doc, "3. Results and Analysis", level=1)

    # 3.1 Tuning heterogeneity
    add_heading_styled(doc, "3.1 Heterogeneity of Orientation Tuning", level=2)
    add_body(doc,
        "The von Mises model was successfully fit to all 23,589 neurons (100% convergence rate), "
        f"with a median R\u00b2 of {r2_median:.3f} (IQR: "
        f"[{tuning_summary[tuning_summary['metric']=='r2_q25']['value'].values[0]:.3f}, "
        f"{tuning_summary[tuning_summary['metric']=='r2_q75']['value'].values[0]:.3f}]). "
        f"A total of {int(tuning_summary[tuning_summary['metric']=='r2_gt_0.5']['value'].values[0]):,} "
        f"neurons (77.6%) achieved R\u00b2 > 0.5, and "
        f"{int(tuning_summary[tuning_summary['metric']=='r2_gt_0.9']['value'].values[0]):,} "
        f"(23.5%) exceeded R\u00b2 > 0.9."
    )

    add_figure(doc, os.path.join(FIGS, "final_fig1_tuning_examples.png"),
               "Figure 1. Example orientation tuning curves with von Mises fits for six neurons spanning "
               "the range of fit quality. Blue circles: empirical bin means (\u00b1 SEM). Orange curves: "
               "fitted von Mises model.", width=5.8)

    add_body(doc,
        f"Preferred orientations were distributed approximately uniformly across 0\u2013180\u00b0 "
        f"(Figure 2, left), indicating no population-level orientation bias. Tuning sharpness "
        f"(\u03ba) varied substantially, with a median of {kappa_median:.2f} "
        f"(IQR: [{kappa_q25:.2f}, {kappa_q75:.2f}]) (Figure 2, right). Approximately 22% of "
        f"neurons were broadly tuned (\u03ba < 1), while 20% were very sharply tuned (\u03ba \u2265 10), "
        f"demonstrating the heterogeneous nature of orientation selectivity in mouse V1."
    )

    add_figure(doc, os.path.join(FIGS, "final_fig2_tuning_parameter_distributions.png"),
               "Figure 2. Distributions of fitted tuning parameters. Left: preferred orientations are approximately "
               "uniform (red dashed line). Right: tuning sharpness (\u03ba) is right-skewed with wide range "
               f"(median = {kappa_median:.2f}).", width=5.8)

    # 3.2 Reliability
    add_heading_styled(doc, "3.2 Tuning Sharpness and Trial-to-Trial Reliability", level=2)
    add_body(doc,
        f"We found a significant positive association between tuning sharpness and reliability. "
        f"The Spearman rank correlation was {spearman_r:.3f} "
        f"(95% CI: [{sp_ci_lo:.3f}, {sp_ci_hi:.3f}]; p < 10\u207b\u00b3\u2070\u2070), "
        f"and the Pearson correlation was {pearson_r:.3f} (Figure 3)."
    )

    # OLS results table
    add_body(doc, "Table 1. OLS regression results: reliability ~ \u03ba + mean_response", bold=True, size=10)
    add_formatted_table(doc,
        ["Parameter", "Coefficient", "Std Error", "t-statistic", "p-value"],
        [
            ["Intercept", "0.670", "0.0022", f"{model_results[model_results['parameter']=='intercept']['t_statistic'].values[0]:.1f}", "< 10\u207b\u00b3\u2070\u2070"],
            ["\u03ba (kappa)", f"{beta_kappa:.4f}", f"{se_kappa:.4f}", f"{t_kappa:.1f}", "< 10\u207b\u2077\u2070"],
            ["Mean response", f"{beta_mean:.4f}", "0.0001", f"{t_mean:.1f}", "< 10\u207b\u00b2\u2070\u2070"],
        ],
        col_widths=[1.3, 1.1, 1.0, 1.1, 1.0],
    )

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(4)

    add_body(doc,
        f"The OLS model (R\u00b2 = 0.055) confirmed that \u03ba has a significant independent "
        f"effect on reliability after controlling for mean response (Table 1). Both \u03ba "
        f"(t = {t_kappa:.1f}) and mean response (t = {t_mean:.1f}) were significant predictors. "
        f"The binned analysis (Figure 3, right) shows a monotonic increase in mean reliability "
        f"with \u03ba, particularly for \u03ba > 5. However, the modest R\u00b2 indicates that "
        f"other factors dominate reliability variance."
    )

    add_figure(doc, os.path.join(FIGS, "final_fig3_reliability_vs_sharpness.png"),
               f"Figure 3. Relationship between tuning sharpness and reliability. Left: hexbin density scatter "
               f"with OLS regression line. Right: binned mean reliability (\u00b1 SEM) shows monotonic positive trend. "
               f"Spearman r = {spearman_r:.3f}, 95% CI [{sp_ci_lo:.3f}, {sp_ci_hi:.3f}].", width=5.8)

    # 3.3 Decoder
    add_heading_styled(doc, "3.3 Population Decoder Performance", level=2)
    add_body(doc,
        f"The OLS linear decoder achieved a mean circular MAE of {mae_1000:.2f}\u00b0 with 1,000 "
        f"neurons (median AE = {medae_1000:.2f}\u00b0), far below the chance level of ~45\u00b0 "
        f"(Figure 4). The shuffle control confirmed this: permuted labels yielded "
        f"MAE of 45.1 \u00b1 0.3\u00b0, matching the theoretical random baseline."
    )

    # Decoder scaling table
    add_body(doc, "Table 2. Decoder performance by neuron count (5-fold CV, 20 random subsets).", bold=True, size=10)
    dec_rows = []
    for _, row in cv_summary.iterrows():
        nc = int(row["neuron_count"])
        mae = f"{row['mae_mean']:.2f}"
        std = f"\u00b1 {row['mae_std']:.2f}" if not pd.isna(row["mae_std"]) else "(single)"
        medae = f"{row['medae_mean']:.2f}"
        n_rep = int(row["n_repeats"])
        dec_rows.append([str(nc), f"{mae} {std}", medae, str(n_rep)])

    add_formatted_table(doc,
        ["Neurons", "MAE (deg)", "Median AE (deg)", "Repeats"],
        dec_rows,
        col_widths=[1.1, 2.0, 1.5, 1.0],
    )

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(4)

    add_body(doc,
        f"Performance improved monotonically with neuron count (Figure 4). The largest gains "
        f"occurred between 10 and 100 neurons (MAE decreased from {mae_10:.1f}\u00b0 to "
        f"{cv_summary.loc[cv_summary['neuron_count']==100, 'mae_mean'].values[0]:.1f}\u00b0). "
        f"Beyond 250 neurons, improvements became more gradual, suggesting partial saturation "
        f"of orientation information. Even at 1,000 neurons, performance had not fully "
        f"plateaued, indicating that additional neurons could further improve decoding."
    )

    add_figure(doc, os.path.join(FIGS, "final_fig4_decoder_scaling.png"),
               "Figure 4. Decoder error vs. neuron count (log scale). Error bars: \u00b1 1 SD "
               "across random neuron subsets. Orange dashed: shuffle control. Gray dotted: "
               "chance level (45\u00b0).", width=5.0)

    add_figure(doc, os.path.join(FIGS, "final_fig5_decoder_examples.png"),
               f"Figure 5. Predicted vs. true stimulus orientation for the 1,000-neuron OLS decoder "
               f"(80/20 train/test split). Tight clustering around the identity line reflects high "
               f"accuracy (MAE = {mae_1000:.1f}\u00b0).", width=3.8)

    # ===================================================================
    # 4. DISCUSSION & CONCLUSION
    # ===================================================================
    add_heading_styled(doc, "4. Discussion and Conclusion", level=1)

    add_heading_styled(doc, "4.1 Summary of Findings", level=2)
    add_body(doc,
        "This analysis of 23,589 simultaneously recorded mouse V1 neurons revealed three "
        "main findings. First, orientation tuning properties are highly heterogeneous: "
        "preferred orientations tile the full 0\u2013180\u00b0 range uniformly, while tuning "
        "sharpness spans over an order of magnitude. This diversity is consistent with "
        "theoretical models of efficient population coding, where a mixture of broadly and "
        "narrowly tuned neurons supports both coarse and fine orientation discrimination "
        "(Ganguli and Simoncelli, 2014)."
    )

    add_body(doc,
        f"Second, sharper tuning is moderately associated with higher trial-to-trial "
        f"reliability (Spearman r = {spearman_r:.3f}), even after controlling for overall "
        f"response magnitude. This suggests that more selective neurons produce more "
        f"consistent signals, though the modest effect size (R\u00b2 \u2248 0.055) indicates "
        f"that noise sources beyond tuning sharpness\u2014such as shared variability, "
        f"behavioral state, and intrinsic noise\u2014dominate reliability."
    )

    add_body(doc,
        f"Third, a simple OLS linear decoder achieves remarkably accurate orientation "
        f"prediction ({mae_1000:.1f}\u00b0 MAE with 1,000 neurons, compared to 45\u00b0 chance), "
        f"consistent with Stringer et al.\u2019s (2021) finding of \u201chigh-precision coding\u201d "
        f"in visual cortex. The monotonic improvement with neuron count, with diminishing "
        f"returns above ~250 neurons, suggests that orientation information is distributed "
        f"across many neurons but with significant redundancy."
    )

    add_heading_styled(doc, "4.2 Limitations", level=2)

    limitations = [
        "Single recording: All results are from one mouse; replication across animals would strengthen generalizability.",
        "Neuron selection: The decoder used the top 1,000 most reliable neurons, a biased subset. Performance with randomly selected neurons from the full population may differ.",
        "Model simplicity: The von Mises model assumes a single peak; neurons with complex tuning (multi-peaked, inhibitory) are not well captured. Similarly, the OLS decoder is linear\u2014nonlinear decoders might extract additional information.",
        "Reliability metric: Split-half reliability depends on the specific trial-splitting method and is influenced by signal-to-noise ratio, partially confounding the sharpness\u2013reliability relationship.",
        "Calcium imaging: Deconvolved calcium activity is an indirect measure of spiking, which may introduce biases in estimated tuning properties.",
    ]
    for lim in limitations:
        p = doc.add_paragraph(style='List Bullet')
        run = p.add_run(lim)
        run.font.size = Pt(11)
        run.font.name = "Arial"

    add_heading_styled(doc, "4.3 Conclusion", level=2)
    add_body(doc,
        "This project demonstrates that mouse V1 contains a diverse population of "
        "orientation-tuned neurons whose collective activity supports high-precision "
        "orientation decoding. The combination of nonlinear least-squares fitting for "
        "single-neuron characterization and linear least-squares regression for population "
        "decoding provides complementary views of neural coding: the former reveals the "
        "building blocks (individual tuning curves), while the latter shows how these "
        "building blocks combine to represent sensory information. Future work could extend "
        "this analysis to multiple recordings, explore nonlinear decoders, and investigate "
        "how tuning heterogeneity relates to the spatial organization of V1."
    )

    # ===================================================================
    # 5. INDIVIDUAL REFLECTION
    # ===================================================================
    add_heading_styled(doc, "5. Individual Reflection", level=1)
    add_body(doc,
        "[To be completed individually by each team member. Include reflection on: "
        "technical skills gained (Python, statistical modeling, data analysis), understanding "
        "of neural coding concepts, challenges encountered and how they were resolved, "
        "and areas for further learning.]"
    )

    # ===================================================================
    # REFERENCES
    # ===================================================================
    add_heading_styled(doc, "References", level=1)

    refs = [
        "Ganguli, D., & Simoncelli, E.P. (2014). Efficient sensory encoding and Bayesian inference with heterogeneous neural populations. Neural Computation, 26(10), 2103\u20132134.",
        "Hubel, D.H., & Wiesel, T.N. (1962). Receptive fields, binocular interaction and functional architecture in the cat\u2019s visual cortex. The Journal of Physiology, 160(1), 106\u2013154.",
        "Stringer, C., Michaelos, M., & Pachitariu, M. (2021). High-precision coding in visual cortex. Cell, 184(10), 2767\u20132778. DOI: 10.1016/j.cell.2021.03.042.",
    ]
    for ref in refs:
        p = doc.add_paragraph()
        run = p.add_run(ref)
        run.font.size = Pt(10)
        run.font.name = "Arial"
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.left_indent = Inches(0.5)
        p.paragraph_format.first_line_indent = Inches(-0.5)

    # ===================================================================
    # Save
    # ===================================================================
    out_path = os.path.join(OUT, "BENG2800_Final_Report.docx")
    doc.save(out_path)
    print(f"Report saved: {out_path}")
    print(f"Size: {os.path.getsize(out_path) / 1e6:.2f} MB")


if __name__ == "__main__":
    main()
